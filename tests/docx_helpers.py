"""Generated Word-document fixtures for shared-pipeline tests."""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

from docx import Document


def make_docx_bytes(blocks: Iterable[tuple[str, object]]) -> bytes:
    """Build a DOCX from ordered paragraph/table blocks in memory."""

    document = Document()
    for kind, value in blocks:
        if kind == "paragraph":
            document.add_paragraph(str(value))
        elif kind == "table":
            rows = [tuple(row) for row in value]  # type: ignore[arg-type]
            if not rows:
                continue
            column_count = max(len(row) for row in rows)
            table = document.add_table(rows=0, cols=column_count)
            for row in rows:
                cells = table.add_row().cells
                for index, cell_value in enumerate(row):
                    cells[index].text = "" if cell_value is None else str(cell_value)
        else:
            raise ValueError(f"Unknown DOCX block type: {kind}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def text_document_blocks(text: str) -> list[tuple[str, object]]:
    """Turn the SDOC text fixture into a title paragraph and field table."""

    lines = [line for line in text.splitlines() if line.strip()]
    title = lines[0] if lines else ""
    rows: list[tuple[str, str]] = []
    for line in lines[1:]:
        if ":" in line:
            label, value = line.split(":", 1)
            rows.append((label, value.strip()))
        else:
            rows.append((line, ""))
    return [("paragraph", title), ("table", rows)]
